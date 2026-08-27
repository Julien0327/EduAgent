"""D2 核心验收脚本。

加载现有种子数据，调用全部 D2 文档服务，生成 4 个文件 + 2 个 HTML 预览，
并重新读取 / 校验生成结果。

运行方式（在仓库根目录执行）：
    .venv/Scripts/python.exe backend/scripts/generate_demo.py
"""
import sys
from pathlib import Path

# 让 app 包可导入（把 backend/ 目录加入 sys.path）
BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.knowledge import loader
from app.services import (
    answer_sheet,
    exam_docx,
    lesson_plan,
    preview,
    report_pdf,
    score_analysis,
)
from app.storage import file_store


def _file_size(path: Path) -> int:
    return path.stat().st_size


def _check_docx(path: Path) -> bool:
    from docx import Document

    doc = Document(str(path))
    return len(doc.paragraphs) > 0


def _check_xlsx(path: Path) -> bool:
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    return wb.active.max_row > 0


def _check_pdf(path: Path) -> bool:
    return path.read_bytes()[:5] == b"%PDF-"


def _check_html(path: Path) -> bool:
    text = path.read_text(encoding="utf-8")
    return len(text.strip()) > 0 and "<html" in text


def main() -> None:
    # 1. 加载种子数据
    bank = loader.load_question_bank()
    answers = loader.load_student_answers()
    title = bank["title"]
    questions = bank["questions"]

    out_dir = file_store.get_outputs_dir("demo")

    # 2. 成绩分析
    result = score_analysis.analyze_scores(answers, bank)
    stats = result["stats"]
    weak_kp = result["weak_knowledge_points"]

    # 3. 教案样例内容（D2 占位，D3 由 LLM 生成）
    lesson_topic = "一元二次方程（配方法）"
    lesson_content = {
        "objectives": [
            "理解配方法解一元二次方程的基本思路，体会转化的数学思想。",
            "掌握用配方法解二次项系数为 1 的一元二次方程。",
        ],
        "key_points": "配方法的步骤与配方技巧。",
        "difficult_points": "如何正确配方并理解（x+m）²=n 的转化过程。",
        "process": [
            ("复习导入", "回顾直接开平方法，引出配方法。"),
            ("新课讲授", "通过例题讲解配方法步骤：移项、配方、开方、求解。"),
            ("例题示范", "例：x² - 6x + 5 = 0 → (x-3)² = 4 → x = 1 或 x = 5。"),
            ("课堂练习", "学生独立完成课本练习题，教师巡视指导。"),
            ("课堂小结", "归纳配方法步骤与注意事项。"),
        ],
        "homework": "完成课后练习第 1、2 题，并预习公式法。",
    }

    # 4. 学情诊断占位（D2 占位文本，D3 由 LLM generate_diagnosis 生成）
    diagnosis = {
        "summary": (
            f"本次课堂小测共 {stats['student_count']} 人参加，班级平均分 {stats['avg']:.1f} 分，"
            f"及格率 {stats['pass_rate'] * 100:.1f}%，优秀率 {stats['excellent_rate'] * 100:.1f}%。"
            "整体掌握情况较好，但部分知识点正确率偏低，需在后续教学中加强。"
        ),
        "diagnosis": (
            "薄弱知识点主要集中在「判别式与根的个数」和「一元二次方程的实际应用」，"
            "说明学生对方程根的情况判断与实际问题建模仍不熟练。"
        ),
        "suggestions": [
            "针对判别式题型进行专项变式训练，强化 Δ=b²-4ac 的符号判断。",
            "应用题重审建模步骤：设未知数 → 列方程 → 求解 → 检验解的合理性。",
            "对 60 分以下学生进行基础补差，对 85 分以上学生布置拓展题，分层施教。",
        ],
    }

    # 5. 生成 4 个文件
    exam_path = Path(
        exam_docx.create_exam_docx(
            title, questions, file_store.build_output_path("demo", "01_试卷.docx")
        )
    )
    answer_path = Path(
        answer_sheet.create_answer_sheet_xlsx(
            title, questions, file_store.build_output_path("demo", "02_答题卡.xlsx")
        )
    )
    report_path = Path(
        report_pdf.create_report_pdf(
            stats, diagnosis, weak_kp, file_store.build_output_path("demo", "03_学情报告.pdf")
        )
    )
    lesson_path = Path(
        lesson_plan.create_lesson_plan_docx(
            lesson_topic, lesson_content, file_store.build_output_path("demo", "04_教案.docx")
        )
    )

    # 6. 生成 2 个 HTML 预览
    exam_html = preview.render_exam_preview(title, questions)
    answer_html = preview.render_answer_sheet_preview(title, questions)
    exam_html_path = file_store.build_output_path("demo", "05_试卷预览.html")
    answer_html_path = file_store.build_output_path("demo", "06_答题卡预览.html")
    exam_html_path.write_text(exam_html, encoding="utf-8")
    answer_html_path.write_text(answer_html, encoding="utf-8")

    # 7. 打印结果
    print("=" * 60)
    print("D2 文档服务验收")
    print("=" * 60)

    print("\n[成绩统计]")
    print(
        f"  班级：{stats['class_name']}，人数：{stats['student_count']}，满分：{stats['full_marks']}"
    )
    print(
        f"  平均分：{stats['avg']:.1f}  中位数：{stats['median']:.1f}  "
        f"最高：{stats['max']}  最低：{stats['min']}  标准差：{stats['std']:.2f}"
    )
    print(
        f"  及格率(≥60)：{stats['pass_rate'] * 100:.1f}%  优秀率(≥85)：{stats['excellent_rate'] * 100:.1f}%"
    )
    print(
        "  分数段分布：",
        {d["label"]: d["count"] for d in stats["score_distribution"]},
    )
    print("  薄弱知识点（按正确率升序）：")
    for kp in weak_kp:
        print(f"    - {kp['name']}（{kp['accuracy'] * 100:.1f}%）")

    print("\n[生成文件]（目录：data/outputs/demo/）")
    files = [exam_path, answer_path, report_path, lesson_path, exam_html_path, answer_html_path]
    for p in files:
        print(f"  {p.name}: {_file_size(p)} bytes")

    print("\n[校验]")
    checks = [
        ("Word 试卷可重新打开", _check_docx(exam_path)),
        ("Word 教案可重新打开", _check_docx(lesson_path)),
        ("Excel 答题卡可重新打开", _check_xlsx(answer_path)),
        ("PDF 报告（%PDF 头）", _check_pdf(report_path)),
        ("HTML 试卷预览", _check_html(exam_html_path)),
        ("HTML 答题卡预览", _check_html(answer_html_path)),
    ]
    all_ok = True
    for name, ok in checks:
        print(f"  {'通过' if ok else '未通过'}：{name}")
        all_ok = all_ok and ok

    print("=" * 60)
    print("全部通过" if all_ok else "存在未通过项", end="")
    print()


if __name__ == "__main__":
    main()
