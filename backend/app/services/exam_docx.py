"""试卷 Word 生成（python-docx）。

题干与答案分离：试卷只含题目（题号 / 题干 / 选项 / 解答留白），不含答案；
标准答案进答题卡（answer_sheet.py）。对应技术方案 §9 / §12。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _setup_font(doc: Document) -> None:
    """设置全局字体（含中文字体 eastAsia），避免中文显示异常。"""
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def create_exam_docx(title: str, questions: list, output_path: str | Path) -> str:
    """根据标题与题目列表生成试卷 docx（不含答案），返回文件路径。"""
    doc = Document()
    _setup_font(doc)

    # 标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    # 信息行（班级 / 姓名 / 得分）
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info.add_run("班级：________　姓名：________　得分：________")
    info_run.font.size = Pt(11)

    # 题目
    for i, q in enumerate(questions, 1):
        # 题号 + 题型 + 分值
        head = doc.add_paragraph()
        head_run = head.add_run(f"{i}.（{q['type']}，{q['score']} 分）")
        head_run.bold = True

        # 题干
        stem = doc.add_paragraph()
        stem.add_run(q["stem"])

        # 选项（选择题）
        if q.get("options"):
            for opt in q["options"]:
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Cm(1)
                opt_p.add_run(opt)

        # 解答题留白
        if q["type"] == "解答题":
            for _ in range(3):
                doc.add_paragraph()

    doc.save(str(output_path))
    return str(output_path)
