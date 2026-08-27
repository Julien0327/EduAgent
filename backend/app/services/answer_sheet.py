"""答题卡 / 评分表 Excel 生成（openpyxl）。

答案进答题卡：本表包含题号、题型、分值、标准答案（含样式）。
对应技术方案 §9 / §12。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def create_answer_sheet_xlsx(
    title: str, questions: list, output_path: str | Path
) -> str:
    """根据标题与题目列表生成答题卡 xlsx（含标准答案），返回文件路径。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "答题卡"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    header_font = Font(bold=True)

    ncols = 4  # 题号 / 题型 / 分值 / 标准答案

    # 标题（合并单元格）
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = center

    # 班级 / 姓名 / 得分 信息行
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    info_cell = ws.cell(
        row=2, column=1, value="班级：____________　姓名：____________　得分：____________"
    )
    info_cell.alignment = center

    # 表头
    headers = ["题号", "题型", "分值", "标准答案"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = header_font
        cell.alignment = center
        cell.border = border
        cell.fill = header_fill

    # 数据行
    for i, q in enumerate(questions, 1):
        row = 3 + i
        values = [i, q["type"], q["score"], q["answer"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.alignment = center
            cell.border = border

    # 列宽
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 30

    wb.save(str(output_path))
    return str(output_path)
