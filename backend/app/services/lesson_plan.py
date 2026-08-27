"""教案 Word 生成（python-docx）。

结构化段落：教学目标 / 教学重难点 / 教学过程 / 课后作业。
对应技术方案 §9 / §12。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _setup_font(doc: Document) -> None:
    """设置全局字体（含中文字体 eastAsia），避免中文显示异常。"""
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def create_lesson_plan_docx(
    topic: str, content: dict, output_path: str | Path
) -> str:
    """根据主题与结构化正文生成教案 docx，返回文件路径。

    content 结构：
      {
        "objectives": [str, ...],
        "key_points": str,
        "difficult_points": str,
        "process": [(step, detail), ...],
        "homework": str,
      }
    """
    doc = Document()
    _setup_font(doc)

    # 标题（主题）
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f"教案：{topic}")
    run.bold = True
    run.font.size = Pt(18)

    # 一、教学目标
    _add_heading(doc, "一、教学目标")
    for obj in content.get("objectives", []):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.add_run(f"• {obj}")

    # 二、教学重点
    _add_heading(doc, "二、教学重点")
    kp = doc.add_paragraph()
    kp.paragraph_format.first_line_indent = Pt(24)
    kp.add_run(content.get("key_points", ""))

    # 三、教学难点
    _add_heading(doc, "三、教学难点")
    dp = doc.add_paragraph()
    dp.paragraph_format.first_line_indent = Pt(24)
    dp.add_run(content.get("difficult_points", ""))

    # 四、教学过程
    _add_heading(doc, "四、教学过程")
    for step, detail in content.get("process", []):
        step_p = doc.add_paragraph()
        step_run = step_p.add_run(step)
        step_run.bold = True
        detail_p = doc.add_paragraph()
        detail_p.paragraph_format.first_line_indent = Pt(24)
        detail_p.add_run(detail)

    # 五、课后作业
    _add_heading(doc, "五、课后作业")
    hw = doc.add_paragraph()
    hw.paragraph_format.first_line_indent = Pt(24)
    hw.add_run(content.get("homework", ""))

    doc.save(str(output_path))
    return str(output_path)
